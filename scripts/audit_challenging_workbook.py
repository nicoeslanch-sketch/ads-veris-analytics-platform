"""Run an explicit local workbook through the production engines for an audit."""

from __future__ import annotations

import argparse
import cProfile
import io
import json
from pathlib import Path
import pstats
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "api"))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--profile", action="store_true")
    parser.add_argument("--metrics", action="store_true")
    parser.add_argument("--compare", type=Path)
    parser.add_argument("--export", type=Path, help="Generate and read back the actual clean XLSX export")
    args = parser.parse_args()
    from app.engine.loader import load_dataframes_with_reports, xlsx_sheet_names
    from app.routes import pipeline
    from app.engine.business import analyze_business_workbook
    from app.engine.multi_sheet import detect_relationships
    from app.support_knowledge import answer_for

    content = args.workbook.read_bytes()
    names = xlsx_sheet_names(content)
    selected = [args.sheet] if args.sheet else names
    started = time.perf_counter()
    frames, _available = load_dataframes_with_reports(args.workbook.name, content, selected)
    audit = {"file": args.workbook.name, "bytes": len(content), "load_seconds": round(time.perf_counter() - started, 3), "sheets": []}
    print(json.dumps({key: value for key, value in audit.items() if key != "sheets"}), flush=True)
    clean_frames, mappings, results = {}, {}, {}
    baseline = json.loads(args.compare.read_text(encoding="utf-8")) if args.compare else None
    previous = {row["sheet"]: row for row in baseline["sheets"]} if baseline else {}
    for name, (frame, report) in frames.items():
        row = {"sheet": name, "rows": len(frame), "columns": list(frame.columns), "exact_duplicates_raw": int(frame.duplicated().sum()), "load_report": report}
        started = time.perf_counter()
        standardized = pipeline._standardize_sync(args.workbook.name, content, name, preloaded=(frame, report))
        row["standardize_seconds"] = round(time.perf_counter() - started, 3)
        row["standardization"] = standardized
        profiler = cProfile.Profile() if args.profile else None
        started = time.perf_counter()
        if profiler:
            profiler.enable()
        try:
            cleaned = pipeline._analyze_cached(args.workbook.name, content, {}, True, sheet=name, preloaded=(frame, report))
            row["cleaning"] = pipeline._public_clean_response(cleaned, args.workbook.name)
            row["clean_rows"] = len(cleaned["_df_limpio"])
            row["dtypes"] = {str(key): str(value) for key, value in cleaned["_df_limpio"].dtypes.items()}
            row["reference_checks"] = {
                "rows_preserved": row["rows"] == row["clean_rows"],
                "duplicates_match_original": row["exact_duplicates_raw"] == cleaned["problemas"]["duplicados"],
                "provenance_preserved": cleaned["_source_rows_limpio"] == list(frame.attrs.get("source_rows", [])),
            }
            if name in previous:
                row["reference_checks"]["unchanged_cleaning_result"] = row["cleaning"] == previous[name]["cleaning"]
            clean_frames[name] = cleaned["_df_limpio"]
            mappings[name] = cleaned["mapeo"]
            results[name] = cleaned
        except Exception as exc:
            row["error"] = f"{type(exc).__name__}: {exc}"
        finally:
            if profiler:
                profiler.disable()
                pstats.Stats(profiler).sort_stats("cumulative").print_stats(25)
        row["clean_seconds"] = round(time.perf_counter() - started, 3)
        if args.metrics and name in results:
            started = time.perf_counter()
            row["metrics"] = pipeline._metrics_from_clean_result(
                args.workbook.name, results[name], None, None, sheet_name=name,
            )
            row["metrics_seconds"] = round(time.perf_counter() - started, 3)
        audit["sheets"].append(row)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(audit, ensure_ascii=False, default=str), encoding="utf-8")
        print(json.dumps({key: row.get(key) for key in ("sheet", "rows", "clean_rows", "exact_duplicates_raw", "standardize_seconds", "clean_seconds", "error")}, ensure_ascii=False), flush=True)

    if args.metrics:
        started = time.perf_counter()
        audit["relationships"] = detect_relationships(clean_frames, mappings)
        audit["relationship_seconds"] = round(time.perf_counter() - started, 3)
        started = time.perf_counter()
        audit["business"] = analyze_business_workbook(clean_frames, mappings, results)
        audit["business_seconds"] = round(time.perf_counter() - started, 3)
        conversations = {
            "Ventas_2024": ["Hola", "cuantossonmisingresostotales", "y mis gastos", "cuanta utilidad tengo", "puedo confiar en esos numeros", "cuantosduplicadoshay", "si los dejo cambia el total", "cual es mi mejor mes", "y el peor", "que canal vende mas", "que categoria lidera", "cuantos clientes tengo", "dame una conclusion", "que informacion falta", "puedo descargar sin borrar duplicados", "por que demora la limpieza", "como conecto las hojas por id", "como importo google sheets", "gracias"],
            "Inventario": ["cuanto stock tengo", "que sucursal tiene mas stock", "y cual tiene menos", "cuanto vale el inventario", "cuantos estan bajo el minimo", "cuantas unidades comprometidas tengo", "cual es la diferencia de conteo", "dame un resumen"],
            "Gastos_Operacionales": ["cuantogaste", "cual es mi gasto neto", "y el iva", "cual es el gasto mas alto", "cual es la categoria principal", "cuanto gaste por categoria", "cual es mi mejor mes", "dame una conclusion"],
            "Metas_Mensuales": ["cuanto es mi meta venta neta", "y la meta margen bruto", "cual es la mediana", "cual es la meta nuevos clientes", "dame un resumen"],
            "Clientes": ["cuantos clientes tengo", "cual es el limite credito promedio", "y el maximo", "que segmento predomina", "dame un resumen"],
        }
        audit["conversations"] = []
        by_name = {row["sheet"]: row for row in audit["sheets"]}
        for name, questions in conversations.items():
            if name not in by_name:
                continue
            history = []
            for question in questions:
                answer = answer_for(question, metrics=by_name[name]["metrics"], history=history[-10:])
                audit["conversations"].append({"sheet": name, "question": question, **answer})
                history.extend([{"role": "user", "content": question}, {"role": "assistant", "content": answer["answer"]}])
        args.output.write_text(json.dumps(audit, ensure_ascii=False, default=str), encoding="utf-8")
        print(json.dumps({"relationship_seconds": audit["relationship_seconds"], "business_seconds": audit["business_seconds"], "conversation_turns": len(audit["conversations"])}), flush=True)

    if args.export:
        import openpyxl
        manifest = {"hojas": [{
            "nombre": name, "procesar": name in results and name != "Parametros",
            "rules": {}, "mapping": {}, "scope": {}, "eliminar_duplicados": False,
        } for name in names]}
        started = time.perf_counter()
        payload, download_name, media_type = pipeline._clean_download_book_sync(
            args.workbook.name, content, manifest, "xlsx", None,
        )
        args.export.parent.mkdir(parents=True, exist_ok=True)
        args.export.write_bytes(payload)
        exported = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
        export_checks = []
        for name, result in results.items():
            if name == "Parametros":
                continue
            sheet = exported[name]
            values = sheet.iter_rows(values_only=True)
            header = next(values)
            exported_rows = list(values)
            expected = result["_df_limpio"]
            check = {
                "sheet": name, "rows": len(exported_rows),
                "row_count_matches": len(exported_rows) == len(expected),
                "headers_match": list(header) == list(expected.columns),
                "exact_duplicates_exported": len(exported_rows) - len(set(exported_rows)),
            }
            if args.metrics and name.startswith("Ventas_"):
                # Independent arithmetic over typed export cells, not the metrics engine.
                source_records = [dict(zip(header, row)) for row in exported_rows]
                eligible = [row for row in source_records
                            if not str(row.get("ID_Documento") or "").upper().startswith("TOTAL")
                            and not str(row.get("Estado") or "").lower().startswith(("anulad", "cancelad", "void"))]
                amounts = [row["Monto Venta"] for row in eligible
                           if isinstance(row.get("Monto Venta"), (float, int))]
                total = round(sum(amounts), 2)
                published = by_name[name]["metrics"]["kpis"]["ingresos_totales"]["valor"]
                check["independent_sales_total"] = total
                check["sales_kpi_matches_export"] = abs(total - published) < 0.01
            export_checks.append(check)
        audit["export"] = {
            "seconds": round(time.perf_counter() - started, 3), "bytes": len(payload),
            "download_name": download_name, "media_type": media_type,
            "sheets": exported.sheetnames, "checks": export_checks,
        }
        exported.close()
        args.output.write_text(json.dumps(audit, ensure_ascii=False, default=str), encoding="utf-8")
        print(json.dumps(audit["export"], ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
