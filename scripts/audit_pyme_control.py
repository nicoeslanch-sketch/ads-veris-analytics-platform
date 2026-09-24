"""Independent, local-only PYME fixture reconciliation; no engine imports/network.

The control workbook is an oracle, never a source of repaired transaction values.
This deliberately reports unrecoverable differences instead of fitting its totals.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from datetime import datetime
import json
from pathlib import Path
import re

import openpyxl


def key(value):
    return str(value).strip().casefold() if value is not None else ""


def number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (float, int, Decimal)):
        return Decimal(str(value))
    text = re.sub(r"[\s$%]", "", str(value)).replace("CLP", "").strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+", text):
        text = text.replace(".", "")
    try:
        result = Decimal(text)
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def discount(value):
    try:
        parsed = Decimal(str(value).replace("%", "").replace(",", ".").strip())
    except InvalidOperation:
        return None
    if not parsed.is_finite():
        return None
    if parsed is None:
        return None
    return parsed / 100 if "%" in str(value) or parsed > 1 else parsed


def records(sheet):
    values = iter(sheet.values)
    columns = next(values)
    return [dict(zip(columns, row)) for row in values if any(v is not None for v in row)]


def cancelled(value):
    return key(value).startswith(("anulad", "cancelad", "void"))


def amount(value):
    return float(value) if value is not None else None


def audit(source: Path, control: Path):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    oracle = openpyxl.load_workbook(control, read_only=True, data_only=True)
    result = {"source": source.name, "control": control.name, "local_only": True,
              "periods": [], "sheets": {}}
    for sheet in workbook:
        values = list(sheet.values)
        result["sheets"][sheet.title] = {"rows_after_first_header": max(len(values) - 1, 0),
            "exact_duplicate_rows": len(values[1:]) - len({tuple(r) for r in values[1:]})}
    product_keys = {key(r["IDProducto"]) for r in records(workbook["Productos"])}
    header_groups = defaultdict(list)
    for name in workbook.sheetnames:
        if name.startswith("Ventas_"):
            for row in records(workbook[name]):
                header_groups[key(row["IDVenta"])].append(row)
    conflicts = {doc for doc, rows in header_groups.items() if len({
        tuple((col, key(v)) for col, v in row.items() if col != "IDVenta") for row in rows
    }) > 1}
    all_sales = {doc: rows[0] for doc, rows in header_groups.items() if doc and doc not in conflicts}
    result["conflicting_headers"] = len(conflicts)
    for name in workbook.sheetnames:
        if not name.startswith("Detalle_20"):
            continue
        counters = Counter()
        totals = Counter()
        identities = {}
        orders = set()
        for row in records(workbook[name]):
            doc = key(row["IDVenta"])
            header = all_sales.get(doc)
            if not header:
                counters["orphan_header"] += 1
                continue
            if cancelled(header["EstadoVenta"]):
                counters["cancelled_lines"] += 1
                continue
            orders.add(doc)
            qty = number(row["Cantidad"])
            price = number(row["PrecioUnitarioNeto_CLP"])
            disc = discount(row["DescuentoPct"])
            declared = number(row["NetoLinea_CLP"])
            cost = number(row["CostoUnitario_CLP"])
            if declared is not None:
                totals["declared_preserving_duplicates"] += declared
            identity = (doc, key(row["Linea"]))
            signature = (key(row["IDProducto"]), qty, price, disc, declared, cost)
            if identity in identities:
                counters["identical_line_copies" if identities[identity] == signature else "conflicting_line_copies"] += 1
                continue
            identities[identity] = signature
            if key(row["IDProducto"]) not in product_keys:
                counters["orphan_product"] += 1
            if declared is not None:
                totals["declared_without_line_copies"] += declared
            valid = qty is not None and qty > 0 and price is not None and price >= 0 and disc is not None and 0 <= disc <= 1
            if not valid:
                counters["invalid_formula_inputs"] += 1
                continue
            computed = (qty * price * (1 - disc)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
            totals["reconstructed_valid_inputs"] += computed
            if declared is not None and abs(computed - declared) > 2:
                counters["net_formula_mismatches"] += 1
            if cost is not None and cost > 0:
                totals["recorded_cost_valid_inputs"] += qty * cost
            else:
                counters["invalid_recorded_cost"] += 1
        result["periods"].append({"period": name.removeprefix("Detalle_"), "orders": len(orders),
                                  "issues": dict(counters), **{k: amount(v) for k, v in totals.items()}})
    expenses = records(workbook["Gastos_Operativos"])
    unique_expenses = {tuple(row.values()): row for row in expenses}
    result["expenses"] = {
        "preserving_duplicates": amount(sum((number(r["MontoNeto_CLP"]) or 0) for r in expenses)),
        "without_exact_copies": amount(sum((number(r["MontoNeto_CLP"]) or 0) for r in unique_expenses.values())),
    }
    receivables = records(workbook["CxC"])
    result["receivables"] = {
        "declared_balance": amount(sum(number(r["Saldo_CLP"]) or 0 for r in receivables)),
        "negative_balances": sum((number(r["Saldo_CLP"]) or 0) < 0 for r in receivables),
        "positive_documents": sum((number(r["Saldo_CLP"]) or 0) > 0 for r in receivables),
    }
    inventory = records(workbook["Inventario_Mensual"])
    dated_inventory = []
    for row in inventory:
        try:
            date = datetime.fromisoformat(str(row["FechaCorte"])).date()
        except (ValueError, TypeError):
            continue
        dated_inventory.append((date, row))
    latest = max((date for date, _ in dated_inventory), default=None)
    snapshot = [row for date, row in dated_inventory if date == latest]
    result["inventory"] = {
        "date": latest.isoformat() if latest else None,
        "rows": len(snapshot), "undated_rows": len(inventory) - len(dated_inventory),
        "stock": amount(sum(number(r["StockUnidades"]) or 0 for r in snapshot)),
        "value": amount(sum((number(r["StockUnidades"]) or 0) * (number(r["CostoPromedio_CLP"]) or 0) for r in snapshot)),
    }
    result["oracle_kpis"] = records(oracle["KPIs_Esperados"])
    result["oracle_periods"] = records(oracle["KPIs_por_Periodo"])
    workbook.close()
    oracle.close()
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("control", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--platform", type=Path, help="Local JSON generated by audit_challenging_workbook.py")
    args = parser.parse_args()
    report = audit(args.source, args.control)
    if args.platform:
        platform = json.loads(args.platform.read_text(encoding="utf-8"))
        published = platform["business"]["estado_resultados"]
        net = sum(row["declared_preserving_duplicates"] for row in report["periods"])
        report["platform_comparison"] = {
            "observed_sales_match": abs(published["ventas_observadas"] - net) < 0.01,
            "expenses_match": abs(published["gastos_operacionales"] - report["expenses"]["preserving_duplicates"]) < 0.01,
            "orphan_headers_match": platform["business"]["alcance"]["filas_sin_cabecera_valida"]
                == sum(row["issues"].get("orphan_header", 0) for row in report["periods"]),
        }
        by_sheet = {row["sheet"]: row.get("metrics", {}) for row in platform["sheets"]}
        inventory = by_sheet.get("Inventario_Mensual", {}).get("analisis_inventario") or {}
        balances = (by_sheet.get("CxC", {}).get("analisis_generico") or {}).get("numericas", [])
        balance = next((row.get("total") for row in balances if row["columna"] == "Saldo_CLP"), None)
        report["platform_comparison"].update({
            "inventory_snapshot_matches": inventory.get("fecha_corte") == report["inventory"]["date"],
            "inventory_stock_matches": inventory.get("stock_total") == report["inventory"]["stock"],
            "inventory_value_matches": inventory.get("valor_inventario") == report["inventory"]["value"],
            "receivables_balance_matches": balance == report["receivables"]["declared_balance"],
        })
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k not in {"sheets", "oracle_kpis", "oracle_periods"}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
