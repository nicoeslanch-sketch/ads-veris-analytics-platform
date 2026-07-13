"""Regresión local del motor con REQ5325, sin copiar datos reales al repo.

Uso:
    python scripts/regresion_req5325.py C:/ruta/REQ5325_....xlsx
    REQ5325_FILE=C:/ruta/archivo.xlsx python scripts/regresion_req5325.py
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "api"))

from app.engine.clean import analyze_and_clean  # noqa: E402
from app.engine.loader import load_dataframe_with_report  # noqa: E402
from app.engine.standardize import (  # noqa: E402
    column_dot3_convention,
    parse_number,
)


EXPECTED = {
    "filas_titulo_omitidas": 1,
    "filas": 14_917,
    "columnas": 16,
    "grupos": 272,
    "filas_involucradas": 865,
    "repeticiones": 593,
    "grupo_maximo": 29,
    "grupos_contiguos": 28,
    "nulos_fisicos": 5_100,
    "placeholders_sin_nombre": 54,
    "espacios_anomalos": 932,
    "mojibake": 4,
    "montos_cero": 669,
    "montos_negativos": 0,
    "montos_invalidos": 0,
    "outliers_iqr": 72,
    "lote_605": 407,
    "lote_605_solapadas": 407,
    "formulas": 0,
}


def _norm_text(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", str(value))
    without_accents = "".join(
        char for char in normalized if not unicodedata.combining(char)
    )
    return " ".join(without_accents.lower().split())


def _has_mojibake(value: object) -> bool:
    text = str(value)
    return any(
        marker in text
        for marker in ("\u00c3", "\u00c2", "\u00ef\u00bf\u00bd", "\ufffd")
    )


def _count_formulas(content: bytes, sheet: str, first_data_row: int, last_data_row: int) -> int:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=True)
    worksheet = workbook[sheet]
    total = 0
    for row in worksheet.iter_rows(min_row=first_data_row, max_row=last_data_row):
        total += sum(cell.data_type == "f" for cell in row)
    workbook.close()
    return total


def run(path: Path) -> dict:
    content = path.read_bytes()
    df, load_report = load_dataframe_with_report(path.name, content)
    safe = analyze_and_clean(df, None, apply=True)
    confirmed = analyze_and_clean(df, None, apply=True, eliminar_duplicados=True)
    duplicate_detail = safe["duplicados_detalle"]

    physical_nulls = sum(
        int(df[column].astype(str).str.strip().eq("").sum()) for column in df.columns
    )
    spacing_changes = sum(
        int(
            (
                df[column].astype(str).map(lambda value: re.sub(r"\s+", " ", value).strip())
                != df[column].astype(str)
            ).sum()
        )
        for column in df.columns
    )
    mojibake = sum(
        sum(_has_mojibake(value) for value in df[column]) for column in df.columns
    )

    client_column = "Raz\u00f3n Social"
    amount_column = "Valor Nominal"
    description_column = "Descripci\u00f3n Lote"
    placeholder_count = int(df[client_column].map(_norm_text).eq("sin nombre").sum())
    convention = column_dot3_convention(df[amount_column])
    amounts = df[amount_column].map(
        lambda value: parse_number(value, dot3_convention=convention)
    )
    lot_605 = df["Lote"].astype(str).str.strip().eq("605")
    empty_description = df[description_column].astype(str).str.strip().eq("")

    source_rows = list(df.attrs["source_rows"])
    facts = {
        "filas_titulo_omitidas": load_report["filas_titulo_omitidas"],
        "filas": len(df),
        "columnas": len(df.columns),
        "grupos": duplicate_detail["grupos"],
        "filas_involucradas": duplicate_detail["filas_involucradas"],
        "repeticiones": duplicate_detail["exactos"],
        "grupo_maximo": duplicate_detail["tamano_maximo_grupo"],
        "grupos_contiguos": duplicate_detail["grupos_contiguos"],
        "nulos_fisicos": physical_nulls,
        "placeholders_sin_nombre": placeholder_count,
        "espacios_anomalos": spacing_changes,
        "mojibake": mojibake,
        "montos_cero": int(amounts.eq(0).sum()),
        "montos_negativos": int(amounts.lt(0).sum()),
        "montos_invalidos": int(amounts.isna().sum()),
        "outliers_iqr": safe["problemas"]["valores_fuera_de_rango"],
        "lote_605": int(lot_605.sum()),
        "lote_605_solapadas": int((lot_605 & amounts.eq(0) & empty_description).sum()),
        "formulas": _count_formulas(
            content,
            load_report["hoja_usada"],
            source_rows[0],
            source_rows[-1],
        ),
    }

    assert facts == EXPECTED, {
        key: {"esperado": EXPECTED[key], "obtenido": value}
        for key, value in facts.items()
        if EXPECTED[key] != value
    }
    assert safe["resumen"]["filas_despues"] == EXPECTED["filas"]
    assert safe["correcciones"]["filas_duplicadas_a_eliminar"] == 0
    assert safe["correcciones"]["filas_duplicadas_eliminadas"] == 0
    assert int(safe["_df_limpio"].duplicated(keep="first").sum()) == 593
    assert confirmed["resumen"]["filas_despues"] == 14_324
    assert confirmed["correcciones"]["filas_duplicadas_eliminadas"] == 593
    assert not {
        "total_problemas", "problemas_totales", "incidencias_totales"
    } & set(safe)

    return {
        **facts,
        "filas_por_defecto": safe["resumen"]["filas_despues"],
        "filas_con_confirmacion": confirmed["resumen"]["filas_despues"],
        # Informativo: cambia cuando evolucionan las reglas; no se fija en EXPECTED.
        "celdas_textuales_reportadas_por_motor": safe["problemas"]["textos_inconsistentes"],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Regresión privada REQ5325")
    parser.add_argument("archivo", nargs="?", help="Ruta al .xlsx real (no se copia)")
    args = parser.parse_args()
    raw_path = args.archivo or os.getenv("REQ5325_FILE")
    if not raw_path:
        parser.error("indica la ruta o define REQ5325_FILE")
    path = Path(raw_path).expanduser().resolve()
    if not path.is_file():
        parser.error(f"no existe el archivo: {path}")

    result = run(path)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    print("Regresión REQ5325: OK. El archivo permaneció fuera del repositorio.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
