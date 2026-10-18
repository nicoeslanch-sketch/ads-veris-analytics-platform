import io

import openpyxl
import pandas as pd

from app.routes import pipeline


def test_single_sheet_download_streams_and_only_converts_the_data_once(monkeypatch):
    frame = pd.DataFrame({'ID': ['001', '002'], 'Importe': ['1.234', '-2.5']})
    original = frame.copy(deep=True)
    result = {'_df_limpio': frame, 'column_types': {'ID': 'texto', 'Importe': 'numero'}}
    monkeypatch.setattr(pipeline, '_analyze_cached', lambda *a, **kw: result)
    monkeypatch.setattr(pipeline, '_build_export_audit', lambda *a: pd.DataFrame({'regla': ['sin_cambio']}))
    monkeypatch.setattr(pipeline, '_export_annotations', lambda *a: ({(0, 'Importe'): True}, {}, []))
    real_export = pipeline.safe_export_dataframe
    calls = []

    def export(*args, **kwargs):
        calls.append(kwargs)
        return real_export(*args, **kwargs)

    monkeypatch.setattr(pipeline, 'safe_export_dataframe', export)
    real_save = pipeline._save_workbook_fast

    def save(wb, output):
        assert wb.write_only
        real_save(wb, output)

    monkeypatch.setattr(pipeline, '_save_workbook_fast', save)
    data, name, media = pipeline._clean_download_sync('datos.csv', b'fixture', {}, 'xlsx')
    assert len(calls) == 1
    pd.testing.assert_frame_equal(frame, original)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert name == 'datos_limpio.xlsx' and 'spreadsheet' in media
    assert wb.sheetnames == ['Datos_limpios', 'Observaciones', 'Auditoria']
    assert list(wb['Datos_limpios'].values) == [('ID', 'Importe'), ('001', 1.234), ('002', -2.5)]
    assert wb['Datos_limpios']['B2'].fill.fgColor.rgb.endswith('FFEB3B')
    assert wb['Datos_limpios'].freeze_panes == 'A2'


def test_untrusted_sheet_names_headers_and_observations_cannot_become_formulas():
    wb = openpyxl.Workbook(write_only=True)
    attack = '=HYPERLINK("https://invalid.example", "click")'
    frame = pd.DataFrame({attack: ['literal']})
    pipeline._write_clean_sheet(wb, 'Datos', frame, {}, {})
    pipeline._write_observations_sheet(wb, [(2, attack, attack, 'revision', attack)])
    pipeline._write_manifest_sheet(wb, [{'hoja': attack, 'estado': 'ok'}])
    output = io.BytesIO()
    wb.save(output)
    restored = openpyxl.load_workbook(io.BytesIO(output.getvalue()), data_only=False)
    cells = [restored['Datos']['A1'], restored['Observaciones']['B2'],
             restored['Observaciones']['C2'], restored['Observaciones']['E2'], restored['Manifest']['A2']]
    assert all(cell.value == attack and cell.data_type == 's' for cell in cells)
