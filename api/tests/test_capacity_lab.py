"""Safety and independent-oracle checks; full stack runs only on manual CI."""

import csv
import importlib.util
import io
from pathlib import Path
from types import SimpleNamespace

import pytest

SPEC = importlib.util.spec_from_file_location('capacity_lab', Path(__file__).resolve().parents[2] / 'scripts/benchmark_durable_local.py')
lab = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(lab)


@pytest.mark.parametrize('url', [
    'https://ads-veris-api.onrender.com', 'https://project.supabase.co',
    'http://localhost:54321', 'http://127.0.0.1.example.org', 'http://2130706433',
    'http://127.0.0.1?target=production', 'http://127.0.0.1#remote',
])
def test_lab_rejects_remote_or_ambiguous_targets(url):
    with pytest.raises(ValueError):
        lab.require_loopback_url(url)


def test_lab_allows_only_explicit_local_services():
    assert lab.require_loopback_url('http://127.0.0.1:54321').port == 54321
    assert lab.require_loopback_url('postgresql://postgres:local@127.0.0.1:54322/postgres', ('postgresql',)).port == 54322


def test_synthetic_oracle_matches_csv_and_accounts_are_distinct():
    one, totals = lab.synthetic_csv(4000, 0)
    two, other_totals = lab.synthetic_csv(4000, 1)
    rows = list(csv.DictReader(io.StringIO(one.decode())))
    assert len(rows) == 4000 and len({r['ID Venta'] for r in rows}) == 4000
    assert sum(totals.values()) == sum(int(r['Monto Venta']) for r in rows)
    assert sum(other_totals.values()) - sum(totals.values()) == 400000
    assert one != two
    for start in ('2026-01-01', '2026-02-12', '2026-02-28'):
        assert sum(v for day, v in totals.items() if day >= start) == sum(int(r['Monto Venta']) for r in rows if r['Fecha'] >= start)


def test_percentiles_use_nearest_rank_in_small_samples():
    assert lab.quantiles([]) == {'count': 0}
    assert lab.quantiles([1, 2, 3]) == {'count': 3, 'median': 2, 'p95': 3, 'max': 3}


def test_workbook_oracle_reconciles_disjoint_periods():
    from openpyxl import load_workbook
    raw, totals, manifest, scope = lab.synthetic_workbook(100, 2)
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    observed, ids = 0, []
    for name in scope['sheets']:
        for row in workbook[name].iter_rows(min_row=2, values_only=True):
            ids.append(row[0].strip())
            observed += int(str(row[5]).replace('$', '').replace('.', '').strip())
    assert observed == sum(totals.values())
    assert len(ids) == len(set(ids)) == 100
    assert len(manifest['hojas']) == len(workbook.sheetnames) == 4
    assert 'Productos' not in scope['sheets']
    workbook.close()


def test_business_capacity_fixture_matches_independent_cost_oracle():
    from app.routes.pipeline import _metrics_multi_sync, _parse_sheet_manifest
    from app.engine.multi_sheet import validate_analysis_scope
    import json

    raw, income, manifest, scope = lab.synthetic_workbook(100, 2, 'business')
    parsed = _parse_sheet_manifest(json.dumps(manifest))
    scope = validate_analysis_scope(scope, [row['nombre'] for row in parsed['hojas']])
    costs = lab.synthetic_cost_totals(100)
    for start in ('2026-01-01', '2026-06-01', '2026-12-01'):
        result = _metrics_multi_sync('synthetic.xlsx', raw, parsed, scope, start, None)
        expected = sum(value for day, value in income.items() if day >= start)
        expected_cost = sum(value for day, value in costs.items() if day >= start)
        assert result['kpis']['ingresos_totales']['valor'] == expected
        assert result['analisis_negocio']['estado_resultados']['ventas_brutas'] == expected
        assert result['analisis_negocio']['estado_resultados']['costo_venta_conocido'] == expected_cost


def test_resource_sampler_sums_all_live_consumers_at_the_same_instant(monkeypatch):
    import psutil
    subject = object.__new__(lab.Lab)
    subject.peak = {}
    subject.aggregate_resources = {'complete_samples': 0, 'incomplete_samples': 0,
                                   'peak_sum_rss_mib': 0, 'max_live_processes': 0}
    subject.processes = [(role, SimpleNamespace(pid=pid, poll=lambda: None), None)
                         for role, pid in [('api', 1), ('worker', 2), ('worker', 3)]]
    memory = {1: 100, 2: 200, 3: 150}
    monkeypatch.setattr(psutil, 'Process', lambda pid: SimpleNamespace(
        memory_info=lambda: SimpleNamespace(rss=memory[pid] * 1024**2)))
    subject.sample_resources_once()
    assert subject.aggregate_resources['peak_sum_rss_mib'] == 450
    memory.update({1: 300, 2: 20, 3: 10})
    subject.sample_resources_once()
    assert subject.peak == {'api': 300, 'worker': 200}
    assert subject.aggregate_resources['peak_sum_rss_mib'] == 450
    assert subject.aggregate_resources['complete_samples'] == 2
    assert subject.aggregate_resources['max_live_processes'] == 3


def test_resource_sampler_does_not_claim_partial_readings_are_complete(monkeypatch):
    import psutil
    subject = object.__new__(lab.Lab)
    subject.peak = {}
    subject.aggregate_resources = {'complete_samples': 0, 'incomplete_samples': 0,
                                   'peak_sum_rss_mib': 0, 'max_live_processes': 0}
    subject.processes = [('api', SimpleNamespace(pid=1, poll=lambda: None), None),
                         ('worker', SimpleNamespace(pid=2, poll=lambda: None), None),
                         ('old_worker', SimpleNamespace(pid=3, poll=lambda: 0), None)]
    def inspect(pid):
        if pid == 2:
            raise psutil.NoSuchProcess(pid)
        assert pid == 1
        return SimpleNamespace(memory_info=lambda: SimpleNamespace(rss=100 * 1024**2))
    monkeypatch.setattr(psutil, 'Process', inspect)
    subject.sample_resources_once()
    assert subject.peak == {'api': 100}
    assert subject.aggregate_resources['complete_samples'] == 0
    assert subject.aggregate_resources['incomplete_samples'] == 1
    assert subject.aggregate_resources['peak_sum_rss_mib'] == 0
